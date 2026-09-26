# -*- coding: utf-8 -*-
"""
测试素材生成器 —— 用于验证解析管道是否正常
============================================
用法：
  python tools\\make_testdata.py            # 生成样例文件
  python tools\\make_testdata.py --clean    # 删除样例文件

生成到：<项目根>\\10_原始归档\\课件\\_样例验证\\
  - 样例课件A.pdf   （3 页：正文页 / 表格页 / 空白页模拟扫描件）
  - 样例课件B.pptx  （2 张幻灯片）
  - 样例教材C.epub  （2 节）

注意：验证完成后请执行 --clean 并删除 20_文本库 中由此产生的记录，
      避免样例数据混入真实知识库。
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
import zipfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = Path(os.environ.get("FIREKB_ROOT") or Path(__file__).resolve().parents[1])
SAMPLE_DIR = ROOT / "10_原始归档" / "课件" / "_样例验证"


def _register_cjk_font() -> str | None:
    """
    注册一个能写中文的字体，返回字体名；都不行则返回 None（退回内置英文字体）。

    为什么要挑字体：这份样例是用来跑通流水线的，生成出来的 PDF **必须能被重新抽取出中文**，
    否则样例只是"看起来对"、其实验不了中文链路。两种方案按可靠性排序：
      ① 系统里的 CJK TrueType 字体（内嵌进 PDF，抽取最可靠）；
      ② ReportLab 自带的 Adobe CID 字体 `STSong-Light`（不内嵌，靠阅读器替换；
         多数情况下也能抽取，但取决于阅读器）。
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for cand in (
        r"C:\Windows\Fonts\msyh.ttc",       # 微软雅黑
        r"C:\Windows\Fonts\simsun.ttc",     # 宋体
        r"C:\Windows\Fonts\simhei.ttf",     # 黑体
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ):
        if not Path(cand).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("FireKBCJK", cand))
            return "FireKBCJK"
        except Exception:
            continue

    try:
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        return "STSong-Light"
    except Exception:
        return None


def make_pdf(path: Path) -> None:
    """
    生成一份样例 PDF（3 页：正文 / 表格样式 / 近乎空白页）。

    用 ReportLab（BSD-3 宽松许可）而不是 PyMuPDF 来写 —— 这样"生成样例"这一步
    也不会把 AGPL 依赖带进项目；生成物本身才是重点，用什么写不重要。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font() or "Helvetica"
    width, height = A4
    c = canvas.Canvas(str(path), pagesize=A4)

    p1_lines = [
        "建筑设计防火规范 学习笔记（样例）",
        "",
        "第一节 防火分区",
        "防火分区是指在建筑内部采用防火墙、楼板及其他防火分隔设施",
        "分隔而成、能在一定时间内防止火灾向同一建筑的其余部分蔓延的",
        "局部空间。",
        "",
        "记忆要点：防火分区的核心是“限定时间内的水平分隔”。",
    ]
    _write_lines(c, p1_lines, font, width, height)
    c.showPage()

    p2_lines = [
        "表 1 不同耐火等级构件的耐火极限（样例数据，非真实条文）",
        "",
        "构件名称          一级          二级          三级",
        "防火墙            3.00h         3.00h         3.00h",
        "承重墙            3.00h         2.50h         2.00h",
        "柱                3.00h         2.50h         2.00h",
        "",
        "注意：上表数字为验证样例，不可用于复习。",
    ]
    _write_lines(c, p2_lines, font, width, height)
    c.showPage()

    c.showPage()          # 第 3 页：几乎空白（模拟扫描件 / 图片页）
    c.save()
    print(f"  生成 {path.name}（字体 {font}）")


def _write_lines(c, lines: list[str], font: str, width: float, height: float,
                 y0: int = 80) -> None:
    """
    逐行写文本。

    ⚠️ 坐标要换算：调用方按"从页面顶部往下数十个 point"给 y（便于读），
    而 PDF/ReportLab 的原点在**左下角** —— 所以写出去之前要 `height - y`。
    （这条换算写错不会报错，只会让整页文字跑到纸外、抽出来是空的。）
    """
    y = y0
    for line in lines:
        if line:
            c.setFont(font, 11)
            c.drawString(56, height - y, line)
        y += 16
        y += 20


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    tb = s1.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(8), Inches(4))
    tf = tb.text_frame
    tf.text = "第二节 安全疏散（样例）"
    for t in [
        "疏散距离：房间内最远点到疏散门的距离",
        "记忆口诀：先看场所，再看层数，最后看距离",
        "老师强调：这一条几乎每年都考，数字必须记准。",
    ]:
        p = tf.add_paragraph()
        p.text = t
        p.font.size = Pt(16)

    s2 = prs.slides.add_slide(blank)
    tb2 = s2.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(8), Inches(4))
    tf2 = tb2.text_frame
    tf2.text = "课堂例题（样例）"
    for t in [
        "题：某建筑高度 24m，求其疏散楼梯间形式。",
        "解：先判断建筑类别，再查对应要求。",
    ]:
        p = tf2.add_paragraph()
        p.text = t
        p.font.size = Pt(16)

    prs.save(str(path))
    print(f"  生成 {path.name}")


def make_epub(path: Path) -> None:
    """手写最小合法 EPUB（EPUB = ZIP + XHTML）。"""
    container = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>"""

    opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="bookid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>消防工程复习样例教材</dc:title>
    <dc:language>zh-CN</dc:language>
    <dc:identifier id="bookid">sample-fire-001</dc:identifier>
  </metadata>
  <manifest>
    <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>
    <item id="c1" href="chap1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="chap2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine toc="ncx">
    <itemref idref="c1"/>
    <itemref idref="c2"/>
  </spine>
</package>"""

    ncx = """<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <head><meta name="dtb:uid" content="sample-fire-001"/></head>
  <docTitle><text>消防工程复习样例教材</text></docTitle>
  <navMap>
    <navPoint id="n1" playOrder="1"><navLabel><text>第一章</text></navLabel><content src="chap1.xhtml"/></navPoint>
    <navPoint id="n2" playOrder="2"><navLabel><text>第二章</text></navLabel><content src="chap2.xhtml"/></navPoint>
  </navMap>
</ncx>"""

    chap1 = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>第一章 火灾基础知识</title></head>
<body>
  <h1>第一章 火灾基础知识</h1>
  <p>火灾是指在时间和空间上失去控制的燃烧所造成的灾害。</p>
  <p>燃烧的必要条件通常概括为可燃物、助燃物与引火源。</p>
  <ul>
    <li>可燃物：能与氧化剂发生燃烧反应的物质</li>
    <li>助燃物：支持燃烧的物质，通常指空气中的氧</li>
    <li>引火源：使可燃物与助燃物发生燃烧反应的能量来源</li>
  </ul>
  <p>记忆要点：三个条件缺一不可，这是判断火灾成因的基础。</p>
</body></html>"""

    chap2 = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>第二章 建筑火灾发展</title></head>
<body>
  <h1>第二章 建筑火灾发展</h1>
  <p>建筑火灾的发展通常划分为初期增长、充分发展、衰减三个阶段。</p>
  <table>
    <tr><th>阶段</th><th>特征</th></tr>
    <tr><td>初期增长</td><td>燃烧面积小，温度上升较慢</td></tr>
    <tr><td>充分发展</td><td>火势蔓延迅速，温度急剧升高</td></tr>
    <tr><td>衰减</td><td>可燃物减少，火势逐渐减弱</td></tr>
  </table>
</body></html>"""

    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        # mimetype 必须是第一个条目且不压缩
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip",
                   compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/toc.ncx", ncx)
        z.writestr("OEBPS/chap1.xhtml", chap1)
        z.writestr("OEBPS/chap2.xhtml", chap2)
    print(f"  生成 {path.name}")


def main() -> int:
    ap = argparse.ArgumentParser(description="生成/清理解析器测试素材")
    ap.add_argument("--clean", action="store_true", help="删除样例目录")
    args = ap.parse_args()

    if args.clean:
        if SAMPLE_DIR.exists():
            shutil.rmtree(SAMPLE_DIR)
            print(f"已删除：{SAMPLE_DIR}")
        else:
            print(f"样例目录不存在：{SAMPLE_DIR}")
        print("\n提醒：如需彻底清除样例痕迹，请删除 <项目根>\\20_文本库\\ 下的")
        print("      pages.jsonl / page_map.json / parse_report.json 后重新解析真实课件。")
        return 0

    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    print(f"生成样例到：{SAMPLE_DIR}")
    n_fail = 0
    try:
        make_pdf(SAMPLE_DIR / "样例课件A.pdf")
    except Exception as exc:
        print(f"  PDF 生成失败：{type(exc).__name__}: {exc}")
        n_fail += 1
    try:
        make_pptx(SAMPLE_DIR / "样例课件B.pptx")
    except Exception as exc:
        print(f"  PPTX 生成失败：{type(exc).__name__}: {exc}")
        n_fail += 1
    try:
        make_epub(SAMPLE_DIR / "样例教材C.epub")
    except Exception as exc:
        print(f"  EPUB 生成失败：{type(exc).__name__}: {exc}")
        n_fail += 1
    if n_fail:
        # 生成失败却退出码 0，会让"照着 README 走"的人以为成功、下一步 parse 才发现没素材。
        # 缺依赖是最常见的原因，所以直接把安装命令给出（这一条正是陌生用户卡住的地方）。
        print(f"\n[FAIL] {n_fail} 个样例没能生成（原因见上）。")
        print("       最常见的原因是没有 reportlab：python -m pip install reportlab")
        return 4
    print("\n下一步：python tools\\parse_docs.py --force")
    return 0


if __name__ == "__main__":
    sys.exit(main())
